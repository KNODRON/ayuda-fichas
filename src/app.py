#!/usr/bin/env python3
"""
app.py - Interfaz para generar texto desde Excel
Requisitos:
  pip install openpyxl
  pip install python-docx   # si vas a generar Word
"""

import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import load_workbook
# from docx import Document  # descomenta si deseas exportar a Word
import os


def get_templates():
    """
    Escanea la carpeta 'templates' y devuelve un diccionario
    {nombre_sin_extension: ruta_completa_al_archivo}
    """
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    templates_dir = os.path.join(base_dir, 'templates')
    files = [f for f in os.listdir(templates_dir) if f.lower().endswith('.xlsx')]
    return {os.path.splitext(f)[0]: os.path.join(templates_dir, f) for f in files}


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Generador de textos desde Excel")
        self.geometry("600x400")

        # Cargar plantillas
        self.templates = get_templates()

        # Selector de plantilla
        ttk.Label(self, text="Selecciona plantilla:").pack(pady=5)
        self.template_cb = ttk.Combobox(self, state="readonly", values=list(self.templates.keys()))
        self.template_cb.pack(pady=5)
        self.template_cb.bind("<<ComboboxSelected>>", self.load_headers)

        # Frame para los campos
        self.form_frame = ttk.Frame(self)
        self.form_frame.pack(fill="x", padx=10, pady=10)
        self.entries = {}

        # Botón de generación
        ttk.Button(self, text="Generar texto", command=self.generate_text).pack(pady=10)
        self.result_txt = tk.Text(self, height=4)
        self.result_txt.pack(fill="both", expand=True, padx=10, pady=5)

    def load_headers(self, event):
        # Limpiar formulario anterior
        for w in self.form_frame.winfo_children():
            w.destroy()
        self.entries.clear()

        # Leer encabezados de la plantilla
        tpl = self.template_cb.get()
        path = self.templates[tpl]
        wb = load_workbook(path, read_only=True)
        sheet = wb.active
        headers = [cell.value for cell in sheet[1]]
        wb.close()

        # Crear un Entry por cada columna
        for idx, header in enumerate(headers, start=1):
            ttk.Label(self.form_frame, text=header).grid(row=idx, column=0, sticky="w", padx=5, pady=2)
            ent = ttk.Entry(self.form_frame)
            ent.grid(row=idx, column=1, sticky="ew", padx=5, pady=2)
            self.entries[header] = ent
        self.form_frame.columnconfigure(1, weight=1)

    def generate_text(self):
        vals = {h: e.get() for h, e in self.entries.items()}
        if not vals:
            messagebox.showwarning("Atención", "Selecciona primero una plantilla y carga los encabezados.")
            return
        try:
            texto = (
                f'» FECHA {vals["Fecha"]}, {vals["Entrada/Salida"]} POR EL TERMINAL {vals["Terminal"]}, '
                f'VÍA {vals["Vía"]}, POR MEDIO DE {vals["Medio"]}, PAÍS {vals["País"]}'
            )
        except KeyError as e:
            messagebox.showerror("Error", f"Falta la columna: {e}")
            return

        # Mostrar en pantalla
        self.result_txt.delete("1.0", tk.END)
        self.result_txt.insert(tk.END, texto)

        # Opcional: exportar a Word
        # doc = Document()
        # doc.add_paragraph(texto)
        # doc.save("salida.docx")


if __name__ == "__main__":
    app = App()
    app.mainloop()
